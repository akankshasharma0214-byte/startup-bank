from datetime import date

import pytest
from openpyxl import load_workbook

from app.billpay import add_vendor, create_bill, schedule_bill_payment
from app.forecast_validate import validate_forecast_workbook
from app.forecast_workbook import build_forecast_workbook
from app.forecasting import build_forecast
from app.invoicing import record_invoice_payment, send_invoice
from app.ledger import get_account
from app.onboarding import fund_account, onboard_company
from app.settlement import run_settlement
from app.simclock import set_sim_date


@pytest.fixture()
def seasoned_company(db):
    set_sim_date(db, date(2026, 1, 5))
    company, admin = onboard_company(
        db, name="Seasoned Co", legal_name="Seasoned Co", ein="0", industry="x", admin_name="A", admin_email="a@a.com", admin_password="x"
    )
    fund_account(db, company_id=company.id, amount_cents=5_000_000, memo="seed")
    vendor = add_vendor(db, company_id=company.id, name="AWS", payment_method="wire")
    software = get_account(db, company.id, "5010")
    revenue = 10_000_00
    for k in range(3):
        m = date(2026, k + 1, 10)
        set_sim_date(db, m)
        inv = send_invoice(db, company_id=company.id, customer_name="Cust", amount_cents=revenue, memo="", due_date=m, created_by=admin)
        record_invoice_payment(db, inv, method="wire", created_by=admin)
        bill = create_bill(
            db,
            company_id=company.id,
            vendor=vendor,
            category=software,
            amount_cents=round(revenue * 0.20),
            memo="",
            due_date=m,
            created_by=admin,
        )
        schedule_bill_payment(db, bill, admin)
        revenue = round(revenue * 1.10)
    set_sim_date(db, date(2026, 4, 1))
    run_settlement(db, company.id)
    return company, admin


def test_workbook_builds_expected_sheets(db, seasoned_company, tmp_path):
    company, admin = seasoned_company
    fc = build_forecast(db, company.id, months=6, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    assert out["sheets"] == ["Historicals", "Assumptions", "IS", "BS", "CF", "Checks"]
    assert len(out["history_cols"]) == 3 and len(out["forecast_cols"]) == 6


def test_workbook_historicals_match_json_history(db, seasoned_company, tmp_path):
    company, admin = seasoned_company
    fc = build_forecast(db, company.id, months=3, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    wh = load_workbook(out["path"])["Historicals"]
    for i, col in enumerate(out["history_cols"]):
        assert wh[f"{col}4"].value == pytest.approx(fc["history"][i]["revenue_cents"] / 100)


def test_forecast_cells_are_formulas_not_values(db, seasoned_company, tmp_path):
    company, admin = seasoned_company
    fc = build_forecast(db, company.id, months=3, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    wb = load_workbook(out["path"])
    for sheet in ("IS", "BS"):
        ws = wb[sheet]
        for col in out["forecast_cols"]:
            for row in range(4, ws.max_row + 1):
                v = ws[f"{col}{row}"].value
                if v is not None:
                    assert isinstance(v, str) and v.startswith("="), (sheet, col, row, v)


def test_recalculated_workbook_passes_all_checks(db, seasoned_company, tmp_path):
    company, admin = seasoned_company
    fc = build_forecast(db, company.id, months=6, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    report = validate_forecast_workbook(out["path"], out["history_cols"], out["forecast_cols"])
    assert report["passed"] is True, report


def test_recalculated_workbook_matches_json_net_income(db, seasoned_company, tmp_path):
    company, admin = seasoned_company
    fc = build_forecast(db, company.id, months=3, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    from pathlib import Path

    from app.forecast_validate import recalculate

    vals, _ = recalculate(Path(out["path"]))
    ws = vals["IS"]
    net_income_row = next(r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value == "Net income")
    for i, col in enumerate(out["forecast_cols"]):
        assert ws[f"{col}{net_income_row}"].value == pytest.approx(fc["income_statement"][i]["net_income_cents"] / 100, abs=0.5)


def test_validate_missing_file():
    assert validate_forecast_workbook("/nonexistent.xlsx", ["B"], ["C"])["passed"] is False


def test_workbook_on_fresh_company_still_validates(db, tmp_path):
    company, admin = onboard_company(
        db, name="Fresh", legal_name="Fresh", ein="0", industry="x", admin_name="A", admin_email="a@a.com", admin_password="x"
    )
    fc = build_forecast(db, company.id, months=3, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    report = validate_forecast_workbook(out["path"], out["history_cols"], out["forecast_cols"])
    assert report["passed"] is True, report


def test_mutated_workbook_detects_broken_balance(db, seasoned_company, tmp_path):
    company, admin = seasoned_company
    fc = build_forecast(db, company.id, months=3, lookback_months=3)
    out = build_forecast_workbook(db, company.id, fc, str(tmp_path / "f.xlsx"))
    wb = load_workbook(out["path"])
    col = out["forecast_cols"][0]
    wb["BS"][f"{col}7"] = 999999  # hardcode equity: breaks both the balance and the no-hardcode check
    wb.save(out["path"])
    report = validate_forecast_workbook(out["path"], out["history_cols"], out["forecast_cols"])
    failed = {c["name"] for c in report["checks"] if not c["passed"]}
    assert "no_hardcoded_forecast_constants" in failed
